# -*- coding: utf-8 -*-
"""第二问模块：utils.py。由原单文件按职责拆分。"""
from __future__ import annotations

import csv
import hashlib
import json
import logging
import math
import sys
from pathlib import Path

import numpy as np
import openpyxl


def require(condition, message):
    if not condition:
        raise ValueError(message)


def clean(value):
    return value.strip() if isinstance(value, str) else value


def number(value, label, positive=False, nonnegative=False):
    require(
        isinstance(value, (int, float, np.number))
        and not isinstance(value, bool),
        f"{label} 应为数值，实际为 {value!r}",
    )
    value = float(value)
    require(math.isfinite(value), f"{label} 不是有限数")
    if positive:
        require(value > 0, f"{label} 必须大于0")
    if nonnegative:
        require(value >= 0, f"{label} 不能为负")
    return value


def integer(value, label, positive=False):
    x = number(value, label, nonnegative=not positive, positive=positive)
    require(abs(x - round(x)) < 1e-9, f"{label} 应为整数")
    return int(round(x))


def sha256(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        while True:
            block = f.read(1024 * 1024)
            if not block:
                break
            h.update(block)
    return h.hexdigest()


def write_csv(path, records, columns=None):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    records = list(records)
    if columns is None:
        columns = list(records[0]) if records else ["message"]
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(records)


def encode_list(values):
    return json.dumps(list(values), ensure_ascii=False)


def init_logging(out):
    logger = logging.getLogger("Q2")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()
    formatter = logging.Formatter("%(asctime)s %(levelname)s %(message)s")
    for handler in [
        logging.StreamHandler(sys.stdout),
        logging.FileHandler(out / "run.log", encoding="utf-8"),
    ]:
        handler.setFormatter(formatter)
        logger.addHandler(handler)
    return logger


def workbook_rows(path, sheet):
    """只读取原始工作簿，不写回；检查Excel错误及无缓存公式。"""
    wf = openpyxl.load_workbook(path, data_only=False)
    wv = openpyxl.load_workbook(path, data_only=True)
    try:
        require(sheet in wf.sheetnames, f"{path.name} 缺少工作表 {sheet}")
        sf, sv = wf[sheet], wv[sheet]
        for row in sf:
            for cell in row:
                require(
                    cell.data_type != "e",
                    f"{path.name}/{sheet}/{cell.coordinate} 存在Excel错误",
                )
                if cell.data_type == "f":
                    require(
                        sv[cell.coordinate].value is not None,
                        f"{path.name}/{sheet}/{cell.coordinate} 公式无缓存值",
                    )
        return [
            tuple(clean(v) for v in row)
            for row in sv.iter_rows(values_only=True)
        ]
    finally:
        wf.close()
        wv.close()


def records_from_header(rows, first_header, expected_headers):
    positions = [
        i for i, row in enumerate(rows)
        if row and row[0] == first_header
        and list(row[:len(expected_headers)]) == expected_headers
    ]
    require(len(positions) == 1, f"表头不唯一或不匹配：{first_header}")
    start = positions[0] + 1
    result = []
    for row in rows[start:]:
        if not row or row[0] is None:
            break
        result.append(dict(zip(expected_headers, row[:len(expected_headers)])))
    return result


def unique_insert(container, key, value, label):
    require(key not in container, f"{label}编号重复：{key}")
    container[key] = value


def raw_metrics(physics, schedule):
    if not schedule:
        return None
    delivery = {}
    for item in schedule:
        c = item["candidate"]
        for b, offset in c.delivery_offset_s.items():
            delivery[b] = item["start_s"] + offset
    return (
        sum(
            physics.data.priority_weight[b]
            * max(0.0, t - physics.data.expected_delivery_s[b])
            for b, t in delivery.items()
        ),
        max(
            item["start_s"] + item["candidate"].duration_s
            for item in schedule
        ),
        sum(item["candidate"].energy_kwh for item in schedule),
        len(schedule),
    )
